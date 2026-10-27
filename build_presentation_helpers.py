# -*- coding: utf-8 -*-
"""
Script to generate the complete 12-slide presentation for Week 4:
"Hệ Thống Báo Khói Cảnh Báo Sớm" - Tầng 2 Network & Connectivity (MQTT) - Nhóm 04
Milestone M2: Connected Device
"""

import sys
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ==============================================================================
# COLOR PALETTE DEFINITION (Navy & Tech Blue Theme)
# ==============================================================================
NAVY_PRIMARY    = RGBColor(31, 78, 120)    # #1F4E78 - Xanh Navy chủ đạo
NAVY_DARK       = RGBColor(15, 43, 70)     # #0F2B46 - Xanh Navy đậm
TECH_BLUE       = RGBColor(46, 117, 182)   # #2E75B6 - Xanh dương công nghệ
TECH_BLUE_LIGHT = RGBColor(222, 235, 247)  # #DEEBF7 - Nền xanh nhạt
TECH_BLUE_ACCENT= RGBColor(0, 164, 239)    # #00A4EF - Xanh dương điểm nhấn
SLATE_DARK      = RGBColor(30, 41, 59)     # #1E293B - Chữ chính đậm
SLATE_MUTED     = RGBColor(100, 116, 139)  # #64748B - Chữ phụ mờ
BG_LIGHT        = RGBColor(248, 250, 252)  # #F8FAFC - Nền slide tổng thể
CARD_BG         = RGBColor(255, 255, 255)  # Trắng card
CARD_BORDER     = RGBColor(205, 220, 235)  # Viền card mềm
ALARM_RED       = RGBColor(192, 0, 0)      # Đỏ cảnh báo nguy hiểm
ALARM_BG        = RGBColor(254, 242, 242)  # Nền đỏ nhạt
SUCCESS_GREEN   = RGBColor(40, 167, 69)    # Xanh lá Pass/Thành công
SUCCESS_BG      = RGBColor(240, 253, 244)  # Nền xanh lá nhạt
WARN_ORANGE     = RGBColor(237, 125, 49)   # Cam lưu ý
CODE_BG         = RGBColor(24, 30, 42)     # Nền dark cho Code JSON
CODE_TEXT       = RGBColor(226, 232, 240)  # Chữ code sáng
WHITE           = RGBColor(255, 255, 255)

FONT_HEADING = "Arial"
FONT_BODY = "Arial"
FONT_CODE = "Consolas"

# ==============================================================================
# HELPER FUNCTIONS FOR SHAPES, TEXT & CARDS
# ==============================================================================
def create_base_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs

def add_slide_header(slide, category_text, title_text):
    """Tạo header chuẩn hóa gồm breadcrumb badge và tiêu đề slide"""
    # Top accent line
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.08))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = TECH_BLUE
    top_bar.line.fill.background()

    # Category Pill / Badge
    cat_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.28), Inches(4.5), Inches(0.32))
    cat_box.fill.solid()
    cat_box.fill.fore_color.rgb = TECH_BLUE_LIGHT
    cat_box.line.color.rgb = TECH_BLUE
    cat_box.line.width = Pt(1)
    tf_cat = cat_box.text_frame
    tf_cat.word_wrap = True
    tf_cat.margin_left = Inches(0.15)
    tf_cat.margin_top = Inches(0.02)
    p_c = tf_cat.paragraphs[0]
    p_c.text = category_text.upper()
    p_c.font.name = FONT_HEADING
    p_c.font.size = Pt(9.5)
    p_c.font.bold = True
    p_c.font.color.rgb = NAVY_PRIMARY

    # Main Slide Title
    title_box = slide.shapes.add_textbox(Inches(0.75), Inches(0.62), Inches(11.8), Inches(0.65))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_top = tf_title.margin_bottom = tf_title.margin_right = 0
    p_t = tf_title.paragraphs[0]
    p_t.text = title_text
    p_t.font.name = FONT_HEADING
    p_t.font.size = Pt(21)
    p_t.font.bold = True
    p_t.font.color.rgb = NAVY_PRIMARY

def add_slide_footer(slide, current_idx, total_slides=12):
    """Tạo footer thanh mảnh hiển thị thông tin nhóm và số trang"""
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(220, 226, 235)
    line.line.fill.background()

    tb_left = slide.shapes.add_textbox(Inches(0.8), Inches(7.08), Inches(8.0), Inches(0.35))
    tf_l = tb_left.text_frame
    tf_l.margin_left = tf_l.margin_top = 0
    p_l = tf_l.paragraphs[0]
    p_l.text = "Đề tài 04: Hệ thống Báo khói Cảnh báo sớm • Tầng 2 Network / Connectivity • Nhóm 04"
    p_l.font.name = FONT_BODY
    p_l.font.size = Pt(8.5)
    p_l.font.color.rgb = SLATE_MUTED

    tb_right = slide.shapes.add_textbox(Inches(10.5), Inches(7.08), Inches(2.0), Inches(0.35))
    tf_r = tb_right.text_frame
    tf_r.margin_right = tf_r.margin_top = 0
    p_r = tf_r.paragraphs[0]
    p_r.alignment = PP_ALIGN.RIGHT
    p_r.text = f"Slide {current_idx:02d} / {total_slides:02d}"
    p_r.font.name = FONT_BODY
    p_r.font.size = Pt(8.5)
    p_r.font.bold = True
    p_r.font.color.rgb = TECH_BLUE

def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER, border_width=1):
    """Vẽ một hộp chứa (card container) với góc bo nhẹ"""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(border_width)
    else:
        card.line.fill.background()
    return card

def add_bullet_point(tf, bold_label, text_content, font_size=11, bold_color=NAVY_PRIMARY, text_color=SLATE_DARK, space_after=6):
    p = tf.add_paragraph()
    p.space_after = Pt(space_after)
    p.line_spacing = 1.15
    if bold_label:
        r1 = p.add_run()
        r1.text = bold_label + " "
        r1.font.name = FONT_BODY
        r1.font.size = Pt(font_size)
        r1.font.bold = True
        r1.font.color.rgb = bold_color
    
    r2 = p.add_run()
    r2.text = text_content
    r2.font.name = FONT_BODY
    r2.font.size = Pt(font_size)
    r2.font.bold = False
    r2.font.color.rgb = text_color
    return p

def add_code_snippet(slide, left, top, width, height, title, code_text):
    """Tạo hộp code dạng terminal IDE hiện đại"""
    card = add_card(slide, left, top, width, height, bg_color=CODE_BG, border_color=RGBColor(50, 60, 80))
    
    # Title bar of code box
    tb_title = slide.shapes.add_textbox(Inches(left + 0.15), Inches(top + 0.08), Inches(width - 0.3), Inches(0.35))
    tf_t = tb_title.text_frame
    tf_t.margin_left = tf_t.margin_top = 0
    p = tf_t.paragraphs[0]
    p.text = f"● ● ●  {title}"
    p.font.name = FONT_HEADING
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(140, 160, 190)

    # Code content
    tb_code = slide.shapes.add_textbox(Inches(left + 0.15), Inches(top + 0.4), Inches(width - 0.3), Inches(height - 0.45))
    tf_c = tb_code.text_frame
    tf_c.word_wrap = True
    tf_c.margin_left = tf_c.margin_top = tf_c.margin_bottom = tf_c.margin_right = 0
    p_code = tf_c.paragraphs[0]
    p_code.text = code_text
    p_code.font.name = FONT_CODE
    p_code.font.size = Pt(9.5)
    p_code.font.color.rgb = CODE_TEXT
    return card

def set_speaker_notes(slide, speaker_name, role_title, notes_text):
    """Gán lời dẫn thuyết trình (Speaker Notes) chi tiết cho từng thành viên"""
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    tf.text = f"【NGƯỜI THUYẾT TRÌNH: {speaker_name.upper()} - {role_title}】\n\n"
    tf.text += f"KỊCH BẢN & LỜI DẪN THỰC TẾ TRƯỚC HỘI ĐỒNG:\n"
    tf.text += notes_text

print("Helper functions initialized successfully.")
